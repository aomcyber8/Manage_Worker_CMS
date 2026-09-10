from __future__ import annotations

import io
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .models import CmsSite, WorkerRecord, WorkerResult, WorkerStatus
from .names import name_key, normalize_name

MAX_UPLOAD = 5 * 1024 * 1024
MAX_ROWS = 50_000
EXPORT_HEADERS = ["วันที่พ้นสภาพ", "ID", "No", "ชื่อ", "E-mail", "ตำแหน่ง", "แผนก", "ผลการค้นหา"]
FIELD_NAMES = {
    "วันที่พ้นสภาพ": "termination_date",
    "id": "id",
    "no": "no",
    "ชื่อ": "name",
    "e-mail": "email",
    "ตำแหน่ง": "position",
    "แผนก": "department",
}
LABELS = {"วันที่พ้นสภาพ": "วันที่พ้นสภาพ", "id": "ID", "no": "No", "ชื่อ": "ชื่อ", "e-mail": "E-mail", "ตำแหน่ง": "ตำแหน่ง", "แผนก": "แผนก"}
DATE_TEXT = re.compile(r"\d{1,2}/\d{1,2}/\d{4}")


def _clean(value: object) -> str:
    return re.sub(r"[\s\u200b\ufeff]+", " ", str(value or "").replace("\xa0", " ")).strip()


def _header(value: object) -> str:
    return _clean(value).casefold()


def _display(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        fmt = str(cell.number_format or "")
        match = re.fullmatch(r"0+", fmt)
        return f"{value:0{len(fmt)}d}" if match else str(value)
    if isinstance(value, float):
        fmt = str(cell.number_format or "")
        if value.is_integer():
            match = re.fullmatch(r"0+", fmt)
            return f"{int(value):0{len(fmt)}d}" if match else str(int(value))
        return str(value)
    return _clean(value)


def parse_excel(raw: bytes, header_row: int = 1, sheet_name: str = "") -> list[WorkerRecord]:
    if not raw:
        raise ValueError("ไฟล์ Excel ว่าง")
    if len(raw) > MAX_UPLOAD:
        raise ValueError("ไฟล์ Excel ต้องไม่เกิน 5 MB")
    if header_row < 1:
        raise ValueError("แถวหัวตารางต้องไม่น้อยกว่า 1")
    try:
        book = load_workbook(io.BytesIO(raw), data_only=True, read_only=False)
    except Exception as exc:
        raise ValueError("อ่านไฟล์ Excel ไม่สำเร็จ") from exc
    try:
        if sheet_name.strip():
            if sheet_name.strip() not in book.sheetnames:
                raise ValueError(f"ไม่พบชีต: {sheet_name.strip()}")
            sheet = book[sheet_name.strip()]
        else:
            sheet = book.active
        if sheet.max_row > MAX_ROWS:
            raise ValueError("รองรับ Excel ไม่เกิน 50,000 แถว")
        found: dict[str, int] = {}
        duplicates: list[str] = []
        for cell in sheet[header_row]:
            normalized = _header(cell.value)
            if normalized not in FIELD_NAMES:
                continue
            if normalized in found:
                duplicates.append(_clean(cell.value))
            else:
                found[normalized] = cell.column
        if duplicates:
            raise ValueError("พบหัวตารางซ้ำ: " + ", ".join(duplicates))
        missing = [LABELS[key] for key in FIELD_NAMES if key not in found]
        if missing:
            raise ValueError("ไม่พบหัวตาราง: " + ", ".join(missing))

        records: list[WorkerRecord] = []
        seen_names: dict[str, int] = {}
        for excel_row in range(header_row + 1, sheet.max_row + 1):
            values = {FIELD_NAMES[key]: _display(sheet.cell(excel_row, column)) for key, column in found.items()}
            if not any(values.values()):
                continue
            missing_values = [label for field, label in (("id", "ID"), ("name", "ชื่อ")) if not values[field]]
            if missing_values:
                raise ValueError(f"แถว {excel_row} ไม่มีข้อมูล: {', '.join(missing_values)}")
            search_name = normalize_name(values["name"])
            if not search_name:
                raise ValueError(f"แถว {excel_row} ไม่มีชื่อสำหรับค้นหา")
            key = name_key(search_name)
            if key in seen_names:
                raise ValueError(f"ชื่อซ้ำหลังตัดคำนำหน้า: {search_name} (แถว {seen_names[key]} และ {excel_row})")
            seen_names[key] = excel_row
            records.append(WorkerRecord(
                excel_row, values["id"], search_name, values["termination_date"], values["id"],
                values["no"], values["name"], values["email"], values["position"], values["department"],
            ))
        if not records:
            raise ValueError("ไม่พบข้อมูล Worker ในไฟล์ Excel")
        return records
    finally:
        book.close()


def _safe_text(value: object) -> str:
    text = str(value or "")
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def _status_label(result: WorkerResult | None) -> str:
    if result is None:
        return "รอดำเนินการ"
    return {
        WorkerStatus.SUCCESS: "Ex-Employee",
        WorkerStatus.SKIPPED: "พบข้อมูล",
        WorkerStatus.NOT_FOUND: "ไม่พบ",
        WorkerStatus.FAILED: "ผิดพลาด",
    }[result.status]


def export_excel(records: list[WorkerRecord], results: list[WorkerResult]) -> bytes:
    if not records:
        raise ValueError("ไม่มีข้อมูลสำหรับ Export")
    result_by_row = {result.row: result for result in results}
    book = Workbook()
    sheet = book.active
    sheet.title = "Workers"
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for column, label in enumerate(EXPORT_HEADERS, 1):
        cell = sheet.cell(1, column, label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    status_fills = {
        "รอดำเนินการ": "D9EAF7", "พบข้อมูล": "DDEBF7", "Ex-Employee": "E2F0D9",
        "ไม่พบ": "FFF2CC", "ผิดพลาด": "F4CCCC",
    }
    for row_index, record in enumerate(records, 2):
        values = [record.termination_date, record.id, record.no, record.name, record.email, record.position, record.department]
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row_index, column)
            if column == 1 and DATE_TEXT.fullmatch(str(value)):
                try:
                    cell.value = datetime.strptime(str(value), "%d/%m/%Y").date()
                    cell.number_format = "dd/mm/yyyy"
                    continue
                except ValueError:
                    pass
            cell.value = _safe_text(value)
            cell.number_format = "@" if column in {2, 3} else "General"
        status = _status_label(result_by_row.get(record.row))
        status_cell = sheet.cell(row_index, 8, status)
        status_cell.fill = PatternFill("solid", fgColor=status_fills[status])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:H{len(records) + 1}"
    for index, width in enumerate((16, 18, 14, 30, 34, 32, 32, 18), 1):
        sheet.column_dimensions[chr(64 + index)].width = width
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def export_filename(source: str, site: CmsSite, now: datetime | None = None) -> str:
    stem = Path(source or "workers").stem
    safe = re.sub(r"[^0-9A-Za-zก-๙._-]+", "_", stem).strip("_") or "workers"
    timestamp = (now or datetime.now()).strftime("%Y%m%d_%H%M%S")
    return f"{safe}_{site.name}_{timestamp}.xlsx"
