from datetime import date
import io

import pytest
from openpyxl import load_workbook

from worker_cms.excel_service import export_excel, export_filename, parse_excel
from worker_cms.models import CmsSite, WorkerResult, WorkerStatus
from conftest import workbook_bytes


ROW = [date(2026, 4, 1), 123, 7, "นายสมชาย ใจดี", "x@example.com", "หัวหน้า", "บุคคล"]


def test_parse_all_columns_and_date():
    item = parse_excel(workbook_bytes(rows=[ROW]))[0]
    assert (item.termination_date, item.id, item.name, item.search_name) == ("01/04/2026", "123", "นายสมชาย ใจดี", "สมชาย ใจดี")


def test_parse_shuffled_headers():
    headers = ["ชื่อ", "แผนก", "ID", "ตำแหน่ง", "E-mail", "วันที่พ้นสภาพ", "No"]
    row = ["นางสาว ก ข", "D", "001", "P", "a@b", "1/4/2026", "09"]
    item = parse_excel(workbook_bytes(headers, [row]))[0]
    assert item.id == "001" and item.department == "D"


def test_missing_header_rejected():
    with pytest.raises(ValueError, match="ไม่พบหัวตาราง: แผนก"): parse_excel(workbook_bytes(headers=["วันที่พ้นสภาพ", "ID", "No", "ชื่อ", "E-mail", "ตำแหน่ง"], rows=[ROW[:-1]]))


def test_duplicate_header_rejected():
    with pytest.raises(ValueError, match="หัวตารางซ้ำ"): parse_excel(workbook_bytes(headers=["ID", "id", "วันที่พ้นสภาพ", "No", "ชื่อ", "E-mail", "ตำแหน่ง", "แผนก"], rows=[]))


@pytest.mark.parametrize(("row", "missing"), [(ROW[:1] + [None] + ROW[2:], "ID"), (ROW[:3] + [None] + ROW[4:], "ชื่อ")])
def test_required_values(row, missing):
    with pytest.raises(ValueError, match=missing): parse_excel(workbook_bytes(rows=[row]))


def test_duplicate_normalized_name_rejected():
    with pytest.raises(ValueError, match="ชื่อซ้ำหลังตัดคำนำหน้า"):
        parse_excel(workbook_bytes(rows=[ROW, [date(2026,4,2), "2", "8", "สมชาย ใจดี", "", "", ""]]))


def test_leading_zero_format_is_preserved():
    from openpyxl import Workbook
    book=Workbook(); sheet=book.active; sheet.append(["วันที่พ้นสภาพ","ID","No","ชื่อ","E-mail","ตำแหน่ง","แผนก"]); sheet.append(["",12,7,"สมชาย ใจดี","","",""])
    sheet["B2"].number_format="00000"; sheet["C2"].number_format="000"
    out=io.BytesIO(); book.save(out); book.close(); item=parse_excel(out.getvalue())[0]
    assert item.id == "00012" and item.no == "007"


def test_export_structure_and_status(record):
    result=WorkerResult(record.row,record.code,record.name,record.search_name,WorkerStatus.SUCCESS,"ok")
    book=load_workbook(io.BytesIO(export_excel([record],[result]))); sheet=book["Workers"]
    assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:H2" and sheet["H2"].value == "Ex-Employee"
    assert sheet["B2"].value == "001" and sheet["C2"].value == "0007"; book.close()


def test_export_formula_injection(record):
    evil=record.__class__(record.row,"=1",record.search_name,record.termination_date,"=1","+2",record.name,"@x",record.position,record.department)
    book=load_workbook(io.BytesIO(export_excel([evil],[]))); sheet=book.active
    assert sheet["B2"].value == "'=1" and sheet["C2"].value == "'+2" and sheet["E2"].value == "'@x"; book.close()


def test_export_requires_data():
    with pytest.raises(ValueError, match="ไม่มีข้อมูล"): export_excel([],[])
def test_export_filename(): assert export_filename("people.xlsx", CmsSite.parse("SENA")).startswith("people_SENA_")
